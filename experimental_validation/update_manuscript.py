"""Surgical edits to a new copy of the supplied manuscript, using measured tables."""
from pathlib import Path
import json
from copy import deepcopy
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT.parent/'manuscript/CausalFed_TGNN_Scientific_Analysis_Edited_Final.docx'
if not SOURCE.exists():
    SOURCE=Path('C:/Users/ehsan/Downloads/CausalFed_TGNN_Scientific_Analysis_Edited_Final.docx')
OUTPUT=ROOT/'CausalFed_TGNN_Experimental_Validation_Complete.docx'
TABLES=ROOT/'results/tables'

def fmt(row,key,digits=3):
    if pd.isna(row[key+'_mean']):return 'PENDING'
    return f'{row[key+"_mean"]:.{digits}f} ± {row[key+"_sd"]:.{digits}f}'

def main():
    manifest=json.loads((ROOT/'RESULTS_MANIFEST.json').read_text());runs=manifest['runs']
    complete=[r for r in runs if r['status']=='COMPLETE'];central=[r for r in complete if r['configuration']['method']=='centralized']
    fed=[r for r in complete if r['configuration']['method']!='centralized' and r['configuration']['attack']=='clean']
    attacks=[r for r in complete if r['configuration']['attack']!='clean']
    summary=pd.read_csv(TABLES/'centralized_and_federated_summary.csv')
    opens=pd.read_csv(TABLES/'Table_E_openset.csv');abl=pd.read_csv(TABLES/'Table_F_ablations.csv');cost=pd.read_csv(TABLES/'Table_I_cost.csv')
    full=summary[(summary.model=='full')&(summary.method=='centralized')&(summary.attack=='clean')]
    assert len(full)==3 and (full.seeds_completed==5).all(),'Do not publish aggregate manuscript before five-seed full-model results exist'
    last=full[full.scale==500000].iloc[0]
    energy=opens[(opens.scale==500000)&(opens.model=='full')&(opens.method=='centralized')&(opens.attack=='clean')&(opens.score=='energy')].iloc[0]
    d=Document(SOURCE);p=list(d.paragraphs);changes=[]
    def replace(i,text):
        old=p[i].text
        prop=deepcopy(p[i].runs[0]._element.rPr) if p[i].runs and p[i].runs[0]._element.rPr is not None else None
        p[i].clear();r=p[i].add_run(text)
        if prop is not None:r._element.insert(0,prop)
        changes.append(dict(paragraph=i,before=old,after=text))
    abstract=(f'Financial fraud detection across institutions combines heterogeneous relations, temporal shift, and fraud mechanisms absent from training. '
       f'We operationalise CausalFed-TGNN as a finite-history heterogeneous graph transformer with invariant/environment branches and independent novelty scores. '
       f'Fifteen immutable synthetic streams at 50,000, 100,000, and 500,000 events use five seeds and chronological 50/20/30 partitions; D/E are excluded from fitting and calibration. '
       f'The recorded screening study contains {len(central)} centralized neural fits, {len(fed)} clean federated simulations, and {len(attacks)} attacked simulations, alongside independently verified historical logistic baselines. '
       f'At 500,000 events, the full variant achieves known-class AUROC {fmt(last,"auroc")}, AP {fmt(last,"ap")}, and F1 {fmt(last,"f1")}; its energy novelty AUROC is {fmt(energy,"auroc")} (mean ± sample SD, five seeds). '
       f'FedAvg, FedProx, SCAFFOLD, and coordinate-median aggregation are compared at 50,000 events with matched attack conditions. '
       f'These are measured screening results, not evidence of causal identification, predictive superiority, or privacy. Fixed daily bank timing, eight-epoch/round budgets, synthetic-only evaluation, and the restricted snapshot architecture limit interpretation. Public-data and continual-adaptation evaluations remain incomplete.')
    replacements={5:abstract,
      15:'The present execution evaluates a finite-history implementation of the proposed graph framework, matched representation ablations, independent novelty scores, and four-bank federated simulations. Results are traceable to saved predictions and checkpoints. This implementation narrows the conceptual architecture: continual adaptation, unrestricted temporal memory, and public-data validation remain pending.',
      16:'The contributions distinguish the broader methodological proposal from the implemented screening study:',
      21:f'An executed evidence package containing {len(central)} centralized neural fits, {len(fed)} clean federated simulations, {len(attacks)} attacked simulations, five-seed summaries, paired ablation differences, per-bank metrics, checkpoints, and independently verified historical logistic evidence.',
      40:'The contribution is an integrated, leakage-controlled evaluation contract and its finite-history implementation. The new neural and federated records permit measured comparisons within this scope; they do not validate every conceptual component or establish predictive superiority, causal invariance, privacy, or real-world generalisation.',
      56:'The implemented prediction graph has one transaction and four bank-local entity snapshots: account, merchant, device, and location. Customer/account identifiers are redundant aliases and are collapsed. Entity-to-transaction context_for edges and reverse uses edges are retained. Each snapshot contains the mean log amount of at most eight previous same-entity events, normalized prior frequency, log elapsed days since the last event, and a history indicator. Snapshots are read before appending the current event. Separate snapshots prevent future entity states from leaking between predictions; earlier unlabelled held-out events may update history prequentially.',
      60:'The implemented encoder uses one PyTorch Geometric HGTConv with width 24 and two attention heads, a residual transaction projection, and ReLU. It reads deterministic finite-history summaries rather than a learned recurrent memory. Reverse-edge outputs do not feed the current transaction in this one-layer readout. Elapsed time is the entity revisit interval, not a variable bank-arrival process; original bank event spacing remains fixed daily. Training-only standardization and event-specific histories prevent held-out information from entering preprocessing.',
      64:'The executed objective is mean fraud cross entropy plus 0.1 times an IRMv1-inspired scale-gradient penalty, 0.01 times supervised cross-bank contrastive loss, and 0.1 times bank-classification loss on the environment branch. Contrastive temperature is 0.2 with at most 128 anchors per shuffled batch. Replay and KL coefficients are zero. IRM and contrastive terms are explicitly disabled for single-bank local batches. The fraud head reads only the invariant branch. Exact formulas and ablation definitions are supplied in EXPERIMENTAL_METHODS.md.',
      66:'Energy is negative logsumexp of the two logits at temperature one. The classifier-only margin comparator is the negative absolute logit difference; prototype distance is minimum Euclidean distance to two training-only class centroids. Every score increases toward unknown, without normalization. Thresholds are the linear 95th percentile of known validation scores, using score ≥ threshold. FPR95 is a descriptive test ROC statistic, not a threshold calibrated on unknown test labels. Amount is retained as an untrained, population-matched novelty comparator.',
      76:'The executed federated experiment simulates four banks in one process with full participation, eight communication rounds, and one local SGD epoch per round at learning rate 0.03. FedAvg uses sample-weighted averaging; FedProx adds a proximal coefficient of 0.01; SCAFFOLD uses local/server control variates; coordinate median averages the middle two coordinates. Checkpoints use lowest pooled validation cross entropy. The server-side validation and different optimizer schedule constrain comparisons with pooled Adam. Single-bank local training disables IRM and cross-bank alignment, so it is not the identical centralized invariant objective.',
      90:'Executed neural comparators are a six-feature Context MLP, a history-matched MLP, homogeneous GraphSAGE, temporal GraphSAGE, static HGT, temporal HGT, and the decomposed full variant. The history-matched MLP receives the same entity summaries as graph models. Static here means no elapsed-time input, while causal histories remain available. Centralized training uses Adam at 0.001, batch size 2048, width 24, and eight epochs. Random forest and boosted-tree candidates remain unexecuted. Historical Amount LR and Context LR evidence is independently verified rather than silently treated as new fits.',
      94:'Note. Label flipping, sign-scaled update poisoning, and Sybil/collusion are executed under the controlled 50k protocol reported below. Backdoor trigger evaluation remains pending. Outcomes establish neither universal attack resistance nor privacy.',
      100:'The canonical RESULTS_MANIFEST.json links each run to its source snapshot, configuration, dataset and split hashes, checkpoint, saved predictions, epoch/round log, environment, and runtime scope. Historical manifests and raw inputs remain unchanged. Known and open-set metrics are independently reconstructed from saved predictions. Source hashes identify the code used at execution, when no Git commit existed for these runs. Later GitHub delivery commits are recorded separately without rewriting historical run provenance.',
      102:'Summaries report mean and sample SD only for five completed distinct seeds at the same scale and configuration. Paired ablation differences retain seed matching. No confidence intervals, p-values, or statistical-significance claims are made. Checkpoints minimize validation cross entropy; classification thresholds maximize validation F1 with highest-threshold tie breaking. Eight epochs or rounds is a fixed screening budget, not evidence of convergence. Model and generator randomness use the same seed and are not disentangled.',
      122:'The original amount novelty diagnostic remains a weak reference on the same late-event population. Trained energy, negative-margin, and prototype-distance scores are now evaluated independently on all test events, with D/E positive and known validation-only thresholds. Table 11 reports their five-seed ranking and operating characteristics. A favorable ranking metric must not be interpreted as an operational detection guarantee; unknown recall and known false-positive rate at the fixed threshold are retained in the machine-readable records.',
      127:'Note. All 15 original streams pass hash, support, chronological separation, unique-ID, and D/E-exclusion checks. Source generator and historical manifests are unchanged. Historical logistic convergence is recorded separately; neural and federated runs complete a fixed screening budget without a convergence claim.',
      129:'6.5 Evidence completion and remaining gates',
      130:f'New recorded evidence comprises {len(central)} centralized neural fits, {len(fed)} clean four-bank federated simulations, and {len(attacks)} attacked simulations. Each completed run retains a reload-tested checkpoint and saved predictions. Graph comparisons, representation ablations, elapsed-time removal on original streams, and trained novelty scores are now evaluated. Continual adaptation, realistic temporal-stress generation, public-data performance, and unseen-bank transfer remain pending. Consequently, the broader project remains experimentally partial.',
      132:'Note. Completion applies only to the specific recorded configuration and scope. The output filename is not a scientific completion claim. Pending modules require separate traceable experiments.',
      135:'OBSERVED: The recorded tables quantify discrimination, novelty detection, ablation differences, and attack effects under one compact training budget. Completion of a graph run does not establish that message passing improves on tabular modeling. The history-matched MLP is needed to distinguish extra input information from architecture. Negative and near-chance outcomes are retained; the protocol does not support a state-of-the-art or significance claim.',
      136:'HYPOTHESIZED: Limited optimization, low predictive signal, class imbalance, and generator-defined mechanism selection may contribute to weak discrimination. These explanations are not established by the present runs. Invariant-loss ablations evaluate temporal generalization within the same four banks; they do not test unseen-bank transfer or identify a causal representation. Changes in branch or relation capacity are also potential confounders and are documented.',
      141:'The evidence is synthetic-only and uses four generator-defined bank environments, five seeds, and eight-epoch/round screening budgets without a hyperparameter search. The implemented one-layer finite-history snapshot architecture is narrower than the conceptual framework. Customer/account redundancy is collapsed; histories are summarized deterministically. Fixed daily bank timing prevents realistic arrival-gap and burst-timing claims. Known-class evaluation excludes D/E and must not be confused with all-fraud screening. Public-data acquisition and authorization remain incomplete.',
      142:'The method does not implement causal discovery, intervention modelling, or causal-effect estimation despite the historical project name. No differential-privacy accountant or secure-aggregation implementation is present, and parameter exchange alone does not establish privacy. Robust aggregation can fail under adaptive or coordinated attacks and should not be described as a universal security defence. Explanation methods can be unstable, incomplete, or sensitive to perturbations and describe model behaviour rather than ground truth.',
      143:'Energy, negative margin, and learned prototype scores are evaluated only under the synthetic holdout protocol. Continual adaptation and explanation-quality evaluation remain pending. The federated simulation disables multi-environment regularization in single-bank local batches and uses pooled validation. Attack assumptions are fixed and limited; sign-flipped updates need not degrade every metric when clean models are weak. Sample SD across five seeds is not a confidence interval or real-world replication guarantee. Conceptual figures remain labelled as such.',
      144:'The expanded package provides verifiable screening evidence while retaining incomplete gates. Longer-budget validation, held-out-bank experiments, variable-arrival stress tests, public data, and continual/explanation evaluations are needed before stronger architectural or deployment claims.',
      146:f'The finite-history CausalFed-TGNN implementation was evaluated through {len(central)} centralized fits, {len(fed)} clean federated simulations, and {len(attacks)} attacked simulations with saved provenance. At 500,000 events its known AUROC is {fmt(last,"auroc")} and energy novelty AUROC is {fmt(energy,"auroc")}. These measured outcomes and matched ablations support a reproducible screening comparison, not a general superiority claim. Synthetic dependence, fixed timing, limited optimization, same-bank evaluation, and missing public/continual validation keep the broader project incomplete.',
      148:'The experimental_validation package supplies source files, pinned dependencies, DATA_AUDIT, RESULTS_MANIFEST.json, individual run records, source snapshots, reload-tested checkpoints, saved late-event predictions, epoch/round logs, generated CSV tables, vector figures, and exact commands in REPRODUCTION.md. Historical ANALYSIS_MANIFEST.json and all raw streams remain unchanged. Source and output hashes permit independent verification. Public-data adapters are schema-tested on artificial fixtures only; no empirical public-data result is included.',
      150:'The executed study uses synthetic records in a single-machine simulation. Client partitioning is an experimental boundary, not a deployed privacy mechanism. Model updates can expose information; no differential privacy or secure aggregation is implemented or claimed. Explanation outputs, if added in future, must describe model behavior rather than causal effects.'}
    for i,text in replacements.items():replace(i,text)
    for i,q in enumerate(p):
        if q.text.strip()=='5.4 Adversarial stress tests':
            q.paragraph_format.page_break_before=True
            q.paragraph_format.keep_with_next=True
            for following in p[i+1:i+5]:
                if not following.text.strip():following.paragraph_format.keep_with_next=True
                elif following.text.startswith('Table 3'):
                    following.paragraph_format.page_break_before=False
                    following.paragraph_format.keep_with_next=True;break
                else:break
    d.tables[2].cell(0,0).text='L = Lsup + λinv Linv + λcon Lcon + λenv Lenv + λreplay Lreplay + λKL LKL'
    gate=d.tables[11]
    for row in list(gate.rows)[1:]:gate._tbl.remove(row._tr)
    for values in [('Scale and provenance','15 immutable streams','PASS'),('Centralized graph models','Five seeds at three scales',f'{len(central)} recorded fits'),
      ('Representation and novelty','Matched ablations and separate scores','EXECUTED within snapshot scope'),
      ('Federated methods','Four methods; 50k; five seeds',f'{len(fed)} clean fits'),('Controlled attacks','Two aggregators; three attacks; five seeds',f'{len(attacks)} fits'),
      ('Public data and continual adaptation','Validated datasets and prequential adaptation','PENDING / BLOCKED'),('Unseen-bank transfer and temporal stress','Separate experiments','PENDING')]:
        for c,v in zip(gate.add_row().cells,values):c.text=v
    anchor=p[133]._p
    def para(text,style=None):
        q=d.add_paragraph(text,style=style);anchor.addprevious(q._p);return q
    def table(caption,headers,rows,note):
        cap=para(caption);cap.paragraph_format.keep_with_next=True
        t=d.add_table(rows=1,cols=len(headers));t.style=d.tables[9].style
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            line=OxmlElement('w:'+edge)
            for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:line.set(qn('w:'+key),value)
            borders.append(line)
        t._tbl.tblPr.append(borders)
        for c,v in zip(t.rows[0].cells,headers):c.text=v
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        for j,row in enumerate(t.rows):
            for k,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF0' if j==0 else ('F5F7F9' if j%2==0 else 'FFFFFF'))
                c._tc.get_or_add_tcPr().append(shade)
                for q in c.paragraphs:
                    q.alignment=WD_ALIGN_PARAGRAPH.LEFT if headers[k] in ['Model','Score','Ablation','Method','Attack'] else WD_ALIGN_PARAGRAPH.CENTER
                    q.paragraph_format.space_after=Pt(3);q.paragraph_format.space_before=Pt(3)
                    for r in q.runs:r.font.size=Pt(8.5);r.bold=(j==0)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        anchor.addprevious(t._tbl);para(note)
    para('6.6 Recorded neural, federated, and attack screening',style=p[129].style)
    para('The following tables are generated from the canonical manifest. COMPLETE means that a defined screening run finished and produced traceable outputs; it does not establish convergence, optimal tuning, or completion of the broader framework.')
    labels={'mlp':'MLP','history_mlp':'History MLP','sage':'GraphSAGE','temporal_sage':'Temporal SAGE','hetero_static':'Static HGT','hgt':'HGT','full':'Full variant',
      'no_inv':'No IRM','no_env':'No environment','no_time':'No elapsed time','no_con':'No contrastive','homogeneous':'Collapsed HGT'}
    s=summary[(summary.method=='centralized')&(summary.attack=='clean')&summary.model.isin(list(labels)[:7])]
    table('Table 10 Executed centralized neural comparison',['Events','Model','AUROC','AP','F1'],
      [[f'{r.scale:,}',labels[r.model],fmt(r,'auroc'),fmt(r,'ap'),fmt(r,'f1')] for _,r in s.iterrows()],
      'Note. Five seeds per row; mean ± sample SD. Positive class A/B/C fraud versus legitimate known test events. F1 uses validation maximum-F1 thresholds. Evidence: completed eight-epoch screening. History MLP controls access to entity summaries; capacity is not identical.')
    q=d.add_paragraph();q.add_run().add_picture(str(ROOT/'results/figures/manuscript_comparison.png'),width=Inches(6.1));anchor.addprevious(q._p)
    para('Fig. 10 Measured known-class AUROC across synthetic scales. Five-seed means and sample SD; fixed eight-epoch screening budget. The dashed line denotes AUROC 0.5.')
    s=opens[(opens.model=='full')&(opens.method=='centralized')&(opens.attack=='clean')]
    table('Table 11 Independent novelty scores for the full variant',['Events','Score','AUROC','AP','FPR95'],
      [[f'{r.scale:,}',r.score.replace('_',' '),fmt(r,'auroc'),fmt(r,'ap'),fmt(r,'fpr95')] for _,r in s.iterrows()],
      'Note. Five seeds; mean ± sample SD; D/E positive among all late test events. Higher score means unknown. Operating thresholds use only the known validation 95th percentile; FPR95 is a descriptive test ROC point. Evidence: completed. Negative margin is the classifier-only scoring ablation; amount is untrained.')
    table('Table 12 Paired representation ablation differences',['Events','Ablation','ΔAUROC','ΔAP','ΔF1'],
      [[f'{r.scale:,}',labels[r.model],fmt(r,'auroc_delta'),fmt(r,'ap_delta'),fmt(r,'f1_delta')] for _,r in abl.iterrows()],
      'Note. Ablation minus full variant, paired by five seeds; mean ± sample SD. A/B/C versus legitimate known test events; each run uses validation maximum-F1 calibration. Evidence: completed screening. Small differences do not establish significance; same-bank results do not demonstrate unseen-bank transfer.')
    s=summary[(summary.method!='centralized')&(summary.attack=='clean')]
    table('Table 13 Clean four-client federated simulation at 50k',['Method','AUROC','AP','F1'],
      [[r.method,fmt(r,'auroc'),fmt(r,'ap'),fmt(r,'f1')] for _,r in s.iterrows()],
      'Note. Five seeds; mean ± sample SD; A/B/C fraud versus legitimate known test events. Pooled validation selects checkpoints and F1 thresholds. Evidence: completed eight-round simulation. Single-bank IRM/alignment are disabled. This is not a private deployment benchmark.')
    s=summary[(summary.method!='centralized')&(summary.attack!='clean')]
    table('Table 14 Controlled attack outcomes at 50k',['Method','Attack','AUROC','AP','F1'],
      [[r.method,r.attack.replace('_',' '),fmt(r,'auroc'),fmt(r,'ap'),fmt(r,'f1')] for _,r in s.iterrows()],
      'Note. Five matched seeds; mean ± sample SD; known A/B/C-versus-legitimate test population. Attack begins at round one: bank 0 flips 50% of labels or sends −5 times its update; Sybil adds two poisoned identities. Pooled validation determines the checkpoint and F1 threshold separately per attacked run. Evidence: completed controlled simulation; no universal robustness claim.')
    s=cost[(cost.method=='centralized')&(cost.attack=='clean')&cost.model.isin(['mlp','sage','hgt','full'])]
    table('Table 15 Recorded computational cost',['Events','Model','Train s','Evaluation s','Parameters'],
      [[f'{r.scale:,}',labels[r.model],fmt(r,'train_seconds',1),fmt(r,'inference_and_prototype_seconds',1),f'{r.parameter_count_mean:,.0f}'] for _,r in s.iterrows()],
      'Note. Five seeds; mean ± sample SD except architecture parameter count. Training includes setup and validation; evaluation includes training prototypes, predictions, and metrics. CPU workers run concurrently. Evidence: measured runtime and deterministic parameter counts. No target class or operating threshold applies to cost. Lifetime memory peaks and tensor-payload communication counts are retained with their scope in CSV records.')
    d.save(OUTPUT)
    (ROOT/'MANUSCRIPT_CHANGELOG.json').write_text(json.dumps(dict(source=str(SOURCE),output=str(OUTPUT),edits=changes,
       added_tables=[10,11,12,13,14,15],study_status='PARTIAL'),indent=2),encoding='utf-8')
    print(OUTPUT)

if __name__=='__main__':main()
